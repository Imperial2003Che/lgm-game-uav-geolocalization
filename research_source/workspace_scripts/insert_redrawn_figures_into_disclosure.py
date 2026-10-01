# -*- coding: utf-8 -*-
from __future__ import annotations

import argparse
import shutil
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


FIGURES = [
    (
        "图5",
        "暗光鲁棒的多模态特征融合与语义门控网络示意图。",
        "本发明暗光鲁棒的多模态特征融合与语义门控网络示意图",
    ),
    (
        "图6",
        "暗光无人机图像的Top-K检索与候选验证示意图。",
        "本发明暗光无人机图像的Top-K检索与候选验证示意图",
    ),
]


def set_run_font(run, size=11, bold=False, east_asia="宋体"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def add_cell_paragraph(cell, text: str):
    p = cell.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(22)
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, size=11, east_asia="宋体")
    return p


def update_appendix_explanation(doc: Document) -> bool:
    updated = False
    for table in doc.tables:
        for row in table.rows:
            if not row.cells:
                continue
            label = "".join(p.text.strip() for p in row.cells[0].paragraphs)
            if "附图说明" not in label:
                continue
            target = row.cells[-1]
            all_text = "\n".join(p.text for p in target.paragraphs)
            for fig_no, desc, _caption in FIGURES:
                if fig_no not in all_text and fig_no.replace(" ", "") not in all_text:
                    add_cell_paragraph(target, f"{fig_no}为{desc}")
                    updated = True
            return updated
    return updated


def add_figure_page(doc: Document, caption: str, image_path: Path, width_inches: float = 6.6):
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(caption)
    set_run_font(run, size=12, bold=True, east_asia="黑体")

    pic_p = doc.add_paragraph()
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.paragraph_format.space_before = Pt(2)
    pic_p.paragraph_format.space_after = Pt(4)
    pic_run = pic_p.add_run()
    pic_run.add_picture(str(image_path), width=Inches(width_inches))


def append_redrawn_figures(doc: Document, fig5: Path, fig6: Path):
    add_figure_page(doc, FIGURES[0][0] + "  " + FIGURES[0][2], fig5)
    add_figure_page(doc, FIGURES[1][0] + "  " + FIGURES[1][2], fig6)


def copy_and_update_md(source_docx: Path, out_docx: Path):
    source_md = source_docx.with_suffix(".md")
    out_md = out_docx.with_suffix(".md")
    if source_md.exists():
        text = source_md.read_text(encoding="utf-8", errors="ignore")
    else:
        text = ""
    addition = (
        "\n\n## 本轮新增重绘附图\n\n"
        "- 图5为本发明暗光鲁棒的多模态特征融合与语义门控网络示意图，展示暗光遮蔽评分、自适应增强、语义修复门控与跨视角检索输出之间的关系。\n"
        "- 图6为本发明暗光无人机图像的Top-K检索与候选验证示意图，展示暗光输入、候选库、语义结构一致性验证、重排序得分和地理位置输出流程。\n"
        "- 本轮附图为重新绘制的本方案示意图，未直接使用参考论文原图。\n"
    )
    out_md.write_text(text + addition, encoding="utf-8")


def append_revision_log(case_dir: Path, out_docx: Path, out_md: Path):
    log_path = case_dir / "交底书修订对话记录.md"
    local_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    utc_time = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
    entry = (
        "\n\n---\n"
        f"记录时间：{local_time}（本地），{utc_time}\n\n"
        "类型：合并迭代\n\n"
        "用户说明摘要：在交底书中加入论文相关图示，但不得直接照抄他人论文原图；如使用参考论文图形结构，应重新绘制并形成有改动的本方案附图。\n\n"
        f"本轮交付文件：{out_md.name}；{out_docx.name}\n\n"
        "合并/修正摘要摘录：本轮在附图说明中新增图5、图6说明，并在附图部分追加两张重新绘制的本发明示意图。新增附图突出云层遮蔽暗光评分、自适应增强、语义修复门控、Top-K检索与候选验证流程；未直接插入参考论文原图。\n"
    )
    if log_path.exists():
        old = log_path.read_text(encoding="utf-8", errors="ignore")
    else:
        old = "# 交底书修订对话记录\n"
    log_path.write_text(old + entry, encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--fig5", type=Path, required=True)
    parser.add_argument("--fig6", type=Path, required=True)
    args = parser.parse_args()

    args.out.parent.mkdir(parents=True, exist_ok=True)
    if args.out.exists():
        args.out.unlink()
    shutil.copy2(args.source, args.out)

    doc = Document(str(args.out))
    update_appendix_explanation(doc)
    append_redrawn_figures(doc, args.fig5, args.fig6)
    doc.save(str(args.out))

    copy_and_update_md(args.source, args.out)
    append_revision_log(args.out.parent, args.out, args.out.with_suffix(".md"))

    print(args.out)
    print(args.out.with_suffix(".md"))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
